"""Run with the bundled Python on a real Windows x64 PC; no production data."""
import asyncio
import json
import pathlib
import platform
import subprocess
import sys
import tempfile

if sys.platform != "win32" or platform.machine().lower() not in {"amd64", "x86_64"}:
    raise SystemExit("This execution test requires a real Windows x64 environment")

import openpyxl
import pypdf
from playwright.sync_api import sync_playwright

with tempfile.TemporaryDirectory(prefix="qianxun-runtime-check-") as folder:
    workbook = openpyxl.Workbook()
    workbook.active.append(["达人", "粉丝", "微信"])
    workbook.active.append(["合成测试", 123, "synthetic-contact"])
    file = pathlib.Path(folder) / "中文资料.xlsx"
    workbook.save(file)
    assert openpyxl.load_workbook(file).active.cell(2, 1).value == "合成测试"
with sync_playwright() as browser:
    assert browser.chromium.name == "chromium"
subprocess.run([sys.executable, "-c", "import asyncio; assert asyncio.run(asyncio.sleep(0, result=3)) == 3"], check=True)
print(json.dumps({"ok": True, "platform": sys.platform, "architecture": platform.machine(),
                  "python": platform.python_version(), "openpyxl": openpyxl.__version__,
                  "pypdf": pypdf.__version__, "playwrightDriverStarted": True,
                  "fullAppAcceptance": False}, ensure_ascii=False))
