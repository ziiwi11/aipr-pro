from __future__ import annotations

import unittest
import tempfile
import json
import io
from pathlib import Path
from unittest.mock import patch
from contextlib import redirect_stdout

import openpyxl

from export_original_format_with_contacts import append_unmatched_rows, build_maps, match_row, main


class OriginalExportTest(unittest.TestCase):
    def match_imported_row(self, values, rows):
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = "名单"
        headers = {"主页身份ID": 1, "抖音主页": 2, "达人昵称": 3, "品牌来源": 4}
        sheet.append(list(headers))
        sheet.append(values)
        return match_row(sheet.title, 2, headers, sheet, *build_maps(rows))

    def test_new_workbook_identity_overrides_old_source_position(self):
        wrong = {"主页身份ID": "old", "源表": "名单", "源行": 2, "达人昵称": "旧达人"}
        correct = {"主页身份ID": "new", "达人昵称": "新达人"}
        self.assertIs(self.match_imported_row(["new", "", "新达人", ""], [wrong, correct]), correct)

    def test_unknown_identity_does_not_inherit_old_contacts(self):
        old = {"主页身份ID": "old", "源表": "名单", "源行": 2, "达人昵称": "同名"}
        self.assertIsNone(self.match_imported_row(["unknown", "", "同名", ""], [old]))

    def test_url_overrides_old_source_position(self):
        old = {"主页身份ID": "old", "源表": "名单", "源行": 2}
        correct = {"抖音主页": "https://douyin.example/new"}
        self.assertIs(self.match_imported_row(["", "https://douyin.example/new?from=search", "", ""], [old, correct]), correct)

    def test_row_position_without_identity_is_not_evidence(self):
        old = {"主页身份ID": "old", "源表": "名单", "源行": 2}
        self.assertIsNone(self.match_imported_row(["", "", "", ""], [old]))

    def test_same_name_without_identity_is_ambiguous(self):
        rows = [{"主页身份ID": "one", "达人昵称": "同名"}, {"主页身份ID": "two", "达人昵称": "同名"}]
        self.assertIsNone(self.match_imported_row(["", "", "同名", ""], rows))

    def test_unique_name_without_identity_can_match(self):
        row = {"主页身份ID": "one", "达人昵称": "唯一", "品牌来源": "品牌"}
        self.assertIs(self.match_imported_row(["", "", "唯一", "品牌"], [row]), row)

    def test_export_completion_is_single_json_event(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "source.xlsx"
            delivery = Path(directory) / "delivery.json"
            book = openpyxl.Workbook()
            book.active.append(["达人昵称", "主页身份ID", "联系方式"])
            book.save(source)
            delivery.write_text(json.dumps({"rows": []}))
            output = io.StringIO()
            with patch("sys.argv", ["export", "--source", str(source), "--delivery", str(delivery), "--out-dir", directory]), redirect_stdout(output):
                main()
            self.assertEqual(len(output.getvalue().splitlines()), 1)
            event = json.loads(output.getvalue())
            self.assertEqual(event["status"], "original_export_finished")
            self.assertTrue(Path(event["output"]).exists())

    def test_macro_workbook_is_rejected_before_read(self):
        with tempfile.TemporaryDirectory() as directory:
            with patch("sys.argv", ["export", "--source", "protected.xlsm", "--delivery", "unused.json", "--out-dir", directory]):
                with self.assertRaisesRegex(ValueError, "仅支持 XLSX"):
                    main()
            self.assertEqual(list(Path(directory).iterdir()), [])

    def test_missing_contact_header_does_not_report_success(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "source.xlsx"
            delivery = Path(directory) / "delivery.json"
            book = openpyxl.Workbook()
            book.active.append(["达人昵称", "无联系方式"])
            book.save(source)
            delivery.write_text(json.dumps({"rows": []}))
            original = source.read_bytes()
            with patch("sys.argv", ["export", "--source", str(source), "--delivery", str(delivery), "--out-dir", directory]):
                with self.assertRaisesRegex(ValueError, "联系方式列"):
                    main()
            self.assertEqual(source.read_bytes(), original)
            self.assertEqual(len(list(Path(directory).glob("*.xlsx"))), 1)

    def test_appends_new_creator_using_original_headers(self) -> None:
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "内容排序池"
        sheet.append(["品牌达人名单"])
        sheet.append([])
        sheet.append([
            "品牌来源", "来源类型", "达人昵称", "主页身份ID", "抖音主页", "综合评分", "等级",
            "精选联盟状态", "抖店精确挂车等级", "联系方式", "联系方式来源", "联系方式状态",
            "建议报价", "报价依据", "数据状态", "品牌/内容证据", "下一步", "微信", "手机号",
        ])
        sheet.row_dimensions[3].height = 44
        row = {
            "达人昵称": "新达人",
            "主页身份ID": "creator-1",
            "抖音主页": "https://douyin.example/1",
            "综合评分": 92,
            "推荐结论": "推荐建联",
            "达人等级": "LV3",
            "明文联系方式": "微信:wx1；手机:13800000000",
            "联系方式来源": "buyin_profile_ui_contact_icon",
            "联系方式状态": "已获取明文",
            "建议报价": 650,
            "推荐理由": "证据完整",
            "验证结论": "主页与内容已验证",
            "内容证据": "真人口播",
            "微信": "wx1",
            "手机号": "13800000000",
        }

        appended = append_unmatched_rows(workbook, [row], "测试品牌", set())

        self.assertEqual(appended, 1)
        self.assertEqual(sheet.cell(4, 3).value, "新达人")
        self.assertEqual(sheet.cell(4, 18).value, "wx1")
        self.assertEqual(sheet.cell(4, 19).value, "13800000000")
        self.assertEqual(sheet.row_dimensions[4].height, 44)


if __name__ == "__main__":
    unittest.main()
