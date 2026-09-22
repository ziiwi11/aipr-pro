from __future__ import annotations

import argparse
import csv
import json
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill


def first(row: dict, *keys: str) -> str:
    for key in keys:
        value = str(row.get(key) or "").strip()
        if value:
            return value
    return ""


def has_contact(row: dict) -> bool:
    return bool(first(row, "buyin_contact_wechat", "cart_contact_wechat", "buyin_contact_phone", "cart_contact_phone"))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--out-dir", required=True)
    args = parser.parse_args()
    payload = json.loads(Path(args.input).read_text(encoding="utf-8"))
    rows = [row for row in payload.get("candidates", []) if isinstance(row, dict) and has_contact(row)]
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    out = Path(args.out_dir)
    out.mkdir(parents=True, exist_ok=True)
    headers = ["序号", "达人昵称", "抖音号", "抖音主页", "精选联盟主页", "达人等级", "类目", "明文联系方式", "微信", "手机号", "联系方式来源", "内容复核", "去重身份"]
    data = []
    for i, row in enumerate(rows, 1):
        data.append([
            i, first(row, "nickname"), first(row, "douyin_id", "unique_id"), first(row, "douyin_homepage"),
            first(row, "buyin_profile_url"), first(row, "author_level", "talent_level"), first(row, "category"),
            first(row, "plain_contact", "buyin_plain_contact", "cart_plain_contact", "buyin_contact_wechat", "cart_contact_wechat", "buyin_contact_phone", "cart_contact_phone"),
            first(row, "buyin_contact_wechat", "cart_contact_wechat"), first(row, "buyin_contact_phone", "cart_contact_phone"),
            first(row, "buyin_contact_source", "cart_contact_source"), first(row, "visual_review_note", "content_review_note"),
            first(row, "creator_id", "author_id", "douyin_id", "unique_id", "buyin_profile_url"),
        ])
    json_path = out / f"aipr_136_contact_delivery_{stamp}.json"
    csv_path = out / f"AIPR当前136个明文联系方式_{stamp}.csv"
    xlsx_path = out / f"AIPR当前136个明文联系方式_{stamp}.xlsx"
    json_path.write_text(json.dumps({"schema_version": "aipr.contact.batch.v1", "created_at": datetime.now().isoformat(), "count": len(rows), "plain_contact_count": len(rows), "wechat_count": sum(bool(x[8]) for x in data), "phone_count": sum(bool(x[9]) for x in data), "candidates": rows}, ensure_ascii=False, indent=2), encoding="utf-8")
    with csv_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(headers)
        writer.writerows(data)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "明文联系方式"
    ws.append(headers)
    for row in data:
        ws.append(row)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="167D68")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths = [8, 20, 20, 42, 42, 10, 18, 26, 22, 18, 24, 32, 32]
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(index)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    wb.save(xlsx_path)
    print(json.dumps({"json": str(json_path), "csv": str(csv_path), "xlsx": str(xlsx_path), "count": len(rows), "wechat": sum(bool(x[8]) for x in data), "phone": sum(bool(x[9]) for x in data)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
