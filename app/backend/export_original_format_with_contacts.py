from __future__ import annotations

import argparse
import copy
import json
import re
from datetime import datetime
from pathlib import Path
from typing import Any

import openpyxl

from console_io import configure_utf8_stdout


configure_utf8_stdout()


def compact(value: Any) -> str:
    return str(value or "").strip()


def norm_url(value: Any) -> str:
    text = compact(value)
    return text.split("?", 1)[0].rstrip("/")


def has_contact(row: dict[str, Any]) -> bool:
    return bool(compact(row.get("明文联系方式") or row.get("微信") or row.get("手机号") or row.get("邮箱")))


def build_maps(rows: list[dict[str, Any]]) -> tuple[dict[str, dict[str, Any]], dict[str, dict[str, Any]], dict[str, dict[str, Any]], dict[tuple[str, int], dict[str, Any]]]:
    by_id: dict[str, dict[str, Any]] = {}
    by_url: dict[str, dict[str, Any]] = {}
    by_name_brand: dict[str, dict[str, Any]] = {}
    by_source: dict[tuple[str, int], dict[str, Any]] = {}
    ambiguous_names: set[str] = set()
    for row in rows:
        if not isinstance(row, dict):
            continue
        keys = [
            compact(row.get("主页身份ID")),
            compact(row.get("精选联盟UID")),
            compact(row.get("验证查询")),
        ]
        for key in keys:
            if key and key not in by_id:
                by_id[key] = row
        url = norm_url(row.get("抖音主页"))
        if url and url not in by_url:
            by_url[url] = row
        name_brand = "|".join([compact(row.get("品牌来源")), compact(row.get("达人昵称"))])
        if compact(row.get("达人昵称")) and name_brand not in ambiguous_names:
            if name_brand in by_name_brand:
                by_name_brand.pop(name_brand)
                ambiguous_names.add(name_brand)
            else:
                by_name_brand[name_brand] = row
        try:
            source_row = int(row.get("源行"))
            source_sheet = compact(row.get("源表"))
            if source_sheet:
                by_source[(source_sheet, source_row)] = row
        except Exception:
            pass
    return by_id, by_url, by_name_brand, by_source


def find_header_row(ws) -> int | None:
    for row_index in range(1, min(ws.max_row or 0, 8) + 1):
        values = [compact(ws.cell(row_index, col).value) for col in range(1, (ws.max_column or 0) + 1)]
        if "达人昵称" in values and ("主页身份ID" in values or "抖音主页" in values):
            return row_index
    return None


def copy_col_style(ws, src_col: int, dst_col: int, max_row: int) -> None:
    ws.column_dimensions[openpyxl.utils.get_column_letter(dst_col)].width = max(
        ws.column_dimensions[openpyxl.utils.get_column_letter(src_col)].width or 12,
        16,
    )
    for row_index in range(1, max_row + 1):
        src = ws.cell(row_index, src_col)
        dst = ws.cell(row_index, dst_col)
        if src.has_style:
            dst._style = copy.copy(src._style)
        if src.number_format:
            dst.number_format = src.number_format
        if src.alignment:
            dst.alignment = copy.copy(src.alignment)
        if src.fill:
            dst.fill = copy.copy(src.fill)
        if src.font:
            dst.font = copy.copy(src.font)
        if src.border:
            dst.border = copy.copy(src.border)


def ensure_column(ws, header_row: int, name: str, style_source_col: int) -> int:
    headers = {compact(ws.cell(header_row, col).value): col for col in range(1, ws.max_column + 1)}
    if name in headers:
        return headers[name]
    col = ws.max_column + 1
    copy_col_style(ws, style_source_col, col, ws.max_row)
    ws.cell(header_row, col).value = name
    return col


def match_row(
    ws_title: str,
    row_index: int,
    headers: dict[str, int],
    ws,
    by_id: dict[str, dict[str, Any]],
    by_url: dict[str, dict[str, Any]],
    by_name_brand: dict[str, dict[str, Any]],
    by_source: dict[tuple[str, int], dict[str, Any]],
) -> dict[str, Any] | None:
    identity = compact(ws.cell(row_index, headers.get("主页身份ID", 0)).value) if headers.get("主页身份ID") else ""
    if identity:
        # A row number belongs to the imported workbook, not to a creator.
        # Explicit identities must never fall back to an old source row/name.
        return by_id.get(identity)
    url = norm_url(ws.cell(row_index, headers.get("抖音主页", 0)).value) if headers.get("抖音主页") else ""
    if url:
        return by_url.get(url)
    name = compact(ws.cell(row_index, headers.get("达人昵称", 0)).value) if headers.get("达人昵称") else ""
    brand = compact(ws.cell(row_index, headers.get("品牌来源", 0)).value) if headers.get("品牌来源") else ""
    if name:
        return by_name_brand.get("|".join([brand, name]))
    return None


def append_unmatched_rows(workbook, rows: list[dict[str, Any]], task_name: str, matched_row_ids: set[int]) -> int:
    target = None
    header_row = None
    for worksheet in workbook.worksheets:
        candidate_header = find_header_row(worksheet)
        if not candidate_header:
            continue
        headers = {compact(worksheet.cell(candidate_header, col).value): col for col in range(1, worksheet.max_column + 1)}
        if "达人昵称" in headers and "联系方式" in headers:
            target = worksheet
            header_row = candidate_header
            break
    if target is None or header_row is None:
        return 0

    headers = {compact(target.cell(header_row, col).value): col for col in range(1, target.max_column + 1)}
    mappings = {
        "品牌来源": lambda row: task_name,
        "来源类型": lambda row: "AI自动采集",
        "达人昵称": lambda row: compact(row.get("达人昵称")),
        "主页身份ID": lambda row: compact(row.get("主页身份ID")),
        "抖音主页": lambda row: compact(row.get("抖音主页")),
        "综合评分": lambda row: row.get("综合评分"),
        "等级": lambda row: "A" if compact(row.get("推荐结论")) == "推荐建联" else "B",
        "精选联盟状态": lambda row: "已验证",
        "抖店精确挂车等级": lambda row: compact(row.get("达人等级")),
        "联系方式": lambda row: compact(row.get("明文联系方式") or row.get("联系方式")),
        "联系方式来源": lambda row: compact(row.get("联系方式来源")),
        "联系方式状态": lambda row: compact(row.get("联系方式状态")),
        "建议报价": lambda row: f"{row.get('建议报价')}元" if row.get("建议报价") not in (None, "") else "",
        "报价依据": lambda row: compact(row.get("推荐理由")),
        "数据状态": lambda row: compact(row.get("验证结论")),
        "品牌/内容证据": lambda row: compact(row.get("内容证据")),
        "下一步": lambda row: "进入机器人建联队列" if compact(row.get("推荐结论")) == "推荐建联" else "人工复核",
        "微信": lambda row: compact(row.get("微信")),
        "手机号": lambda row: compact(row.get("手机号")),
    }

    appended = 0
    template_row = target.max_row
    template_height = target.row_dimensions[template_row].height
    for row in rows:
        if id(row) in matched_row_ids:
            continue
        row_index = target.max_row + 1
        style_row = max(header_row + 1, template_row)
        for col in range(1, target.max_column + 1):
            source_cell = target.cell(style_row, col)
            destination = target.cell(row_index, col)
            if source_cell.has_style:
                destination._style = copy.copy(source_cell._style)
            destination.number_format = source_cell.number_format
            destination.alignment = copy.copy(source_cell.alignment)
        target.row_dimensions[row_index].height = template_height
        for header, getter in mappings.items():
            if header in headers:
                target.cell(row_index, headers[header]).value = getter(row)
        appended += 1
    return appended


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", required=True)
    parser.add_argument("--delivery", required=True)
    parser.add_argument("--out-dir", default="output")
    parser.add_argument("--task-name", default="品牌任务")
    args = parser.parse_args()
    out_dir = Path(args.out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    source = Path(args.source)
    if source.suffix.lower() != ".xlsx":
        raise ValueError("原格式导出仅支持 XLSX；请先另存为 XLSX，避免丢失宏")
    delivery = Path(args.delivery)
    data = json.loads(delivery.read_text(encoding="utf-8"))
    rows = [row for row in data.get("rows") or [] if isinstance(row, dict)]
    by_id, by_url, by_name_brand, by_source = build_maps(rows)

    workbook = openpyxl.load_workbook(source)
    if not any((header := find_header_row(ws)) and "联系方式" in [compact(ws.cell(header, col).value) for col in range(1, ws.max_column + 1)] for ws in workbook.worksheets):
        raise ValueError("原始表格需包含达人昵称、主页身份ID或抖音主页中的身份列，以及联系方式列；原文件未修改")
    touched_rows = 0
    contact_rows = 0
    matched_row_ids: set[int] = set()
    for ws in workbook.worksheets:
        header_row = find_header_row(ws)
        if not header_row:
            continue
        headers = {compact(ws.cell(header_row, col).value): col for col in range(1, ws.max_column + 1)}
        if "联系方式" not in headers:
            continue
        contact_col = headers["联系方式"]
        source_col = headers.get("联系方式来源")
        status_col = headers.get("联系方式状态")
        wechat_col = ensure_column(ws, header_row, "微信", contact_col)
        phone_col = ensure_column(ws, header_row, "手机号", contact_col)

        for row_index in range(header_row + 1, ws.max_row + 1):
            matched = match_row(ws.title, row_index, headers, ws, by_id, by_url, by_name_brand, by_source)
            if not matched:
                continue
            matched_row_ids.add(id(matched))
            touched_rows += 1
            if has_contact(matched):
                contact_rows += 1
                ws.cell(row_index, contact_col).value = compact(matched.get("明文联系方式") or matched.get("联系方式"))
                ws.cell(row_index, wechat_col).value = compact(matched.get("微信"))
                ws.cell(row_index, phone_col).value = compact(matched.get("手机号"))
                if source_col:
                    ws.cell(row_index, source_col).value = compact(matched.get("联系方式来源") or "buyin_profile_ui_contact_icon")
                if status_col:
                    ws.cell(row_index, status_col).value = "已获取明文"

    appended_rows = append_unmatched_rows(workbook, rows, args.task_name, matched_row_ids)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    safe_name = re.sub(r'[\\/:*?"<>|]+', "_", args.task_name).strip() or "品牌任务"
    output = out_dir / f"{safe_name}_原格式补联系方式_{stamp}.xlsx"
    workbook.save(output)
    print(
        json.dumps(
            {
                "status": "original_export_finished",
                "output": str(output),
                "delivery": str(delivery),
                "touched_rows": touched_rows,
                "contact_rows": contact_rows,
                "appended_rows": appended_rows,
                "plain_contact_count": data.get("plain_contact_count"),
                "wechat_contact_count": data.get("wechat_contact_count"),
            },
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
