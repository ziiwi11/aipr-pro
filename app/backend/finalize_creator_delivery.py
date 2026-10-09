from __future__ import annotations

import argparse
import json
import os
import subprocess
import sys
from datetime import datetime
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from atomic_json_io import atomic_write_json
from audit_strict_contact_highwater import _assert_strict_invariants
from console_io import configure_utf8_stdout
from creator_delivery_contract import build_delivery, select_delivery_candidates, creator_identity_values, contact_identity_values


configure_utf8_stdout()


def emit(payload: dict[str, Any]) -> None:
    print(json.dumps(payload, ensure_ascii=False), flush=True)


def write_standard_workbook(rows: list[dict[str, Any]], target: Path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "最终达人名单"
    headers = list(rows[0].keys()) if rows else ["序号", "达人昵称", "推荐结论", "明文联系方式", "微信", "手机号"]
    sheet.append(headers)
    for row in rows:
        sheet.append([row.get(header, "") for header in headers])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="167D68")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        values = [str(cell.value or "") for cell in column[:80]]
        width = min(42, max(10, max((len(value) for value in values), default=10) + 2))
        sheet.column_dimensions[column[0].column_letter].width = width
        for cell in column:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    workbook.save(target)


def merge_robot_queue(target: Path, records: list[dict[str, Any]]) -> None:
    existing: list[dict[str, Any]] = []
    if target.exists():
        for line in target.read_text(encoding="utf-8").splitlines():
            if line.strip():
                try:
                    existing.append(json.loads(line))
                except json.JSONDecodeError:
                    continue
    def record_key(row: dict[str, Any]) -> tuple[str, str]:
        return (
            str((row.get("task") or {}).get("id") or row.get("taskId") or ""),
            str((row.get("creator") or {}).get("id") or row.get("creatorId") or ""),
        )

    merged = {record_key(row): row for row in existing}
    for row in records:
        merged[record_key(row)] = row
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text("".join(json.dumps(row, ensure_ascii=False) + "\n" for row in merged.values()), encoding="utf-8")


def _manifest_roi(roi: Any) -> dict[str, Any]:
    """从 ROI 报告中提取交接清单需要的摘要字段。"""
    if not isinstance(roi, dict) or roi.get("error"):
        return {"available": False, "reason": (roi or {}).get("error", "not_computed")}
    funnel = roi.get("funnel") or {}
    contacts = roi.get("contacts") or {}
    return {
        "available": True,
        "funnel": {
            "candidates": funnel.get("candidates", 0),
            "verified": funnel.get("verified", 0),
            "qualified": funnel.get("qualified", 0),
            "revealed": funnel.get("revealed", 0),
            "delivered": funnel.get("delivered", 0),
        },
        "rates": funnel.get("rates") or {},
        "contacts": {
            "withContact": contacts.get("with_contact", 0),
            "phone": contacts.get("phone", 0),
            "wechat": contacts.get("wechat", 0),
            "contactRate": contacts.get("contact_rate", "0.0"),
        },
        "risk": roi.get("risk") or {},
    }


def batch_new_candidates(rows: list[dict], baseline: list[dict]) -> list[dict]:
    if not baseline or any(not creator_identity_values(row) for row in baseline):
        raise ValueError("batch_baseline_requires_strong_identities")
    seen = set().union(*(creator_identity_values(row) for row in baseline))
    return [row for row in rows if not creator_identity_values(row) & seen]


def contact_completion_summary(rows: list[dict], target_count: int, baseline_path: Path) -> dict:
    checked = [row for row in rows if row.get("ui_contact_channels_checked_at")]
    reasons: dict[str, int] = {}
    for row in rows:
        state = row.get("ui_contact_channels_status") or "not_checked"
        if any(marker in str(row.get("ui_contact_message") or "") for marker in ("主推类目不属于店铺类目", "本周仅支持")):
            state = "category_not_matched"
        reasons[state] = reasons.get(state, 0) + 1
    baseline_rows = None
    baseline_status = "missing"
    if baseline_path.is_file():
        try:
            payload = json.loads(baseline_path.read_text(encoding="utf-8"))
            baseline_rows = payload.get("candidates")
            if not isinstance(baseline_rows, list) or any(not isinstance(row, dict) or not creator_identity_values(row) for row in baseline_rows):
                raise ValueError("invalid contact baseline")
            baseline_status = "available"
        except (ValueError, OSError, AttributeError):
            baseline_rows = None
            baseline_status = "unreadable"
    def has_wechat(row: dict) -> bool:
        return bool(row.get("buyin_contact_wechat") or row.get("cart_contact_wechat"))
    new_wechat = None
    if baseline_rows is not None:
        old_with_wechat = set().union(*(creator_identity_values(row) for row in baseline_rows if has_wechat(row)))
        new_wechat = sum(has_wechat(row) and not creator_identity_values(row) & old_with_wechat for row in checked)
    return {
        "checked_count": len(checked), "target_count": target_count, "status_counts": reasons,
        "new_wechat_count": new_wechat,
        "supplement_duplicate_channels_rejected": sum(len(row.get("ui_contact_supplement_duplicate_channels") or []) for row in rows),
        "baseline_available": baseline_rows is not None,
        "baseline_status": baseline_status,
        "baseline_path": str(baseline_path) if baseline_rows is not None else "",
        "baseline_note": "" if baseline_rows is not None else "缺少可核对的渠道补全前快照，新增微信数量未知；不将现有微信算作补全新增",
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--strategy", required=True)
    parser.add_argument("--out-dir", required=True)
    parser.add_argument("--task-id", default="brand-task")
    parser.add_argument("--task-name", default="品牌达人任务")
    parser.add_argument("--original-workbook", default="")
    parser.add_argument("--robot-queue", default="")
    parser.add_argument("--contacts-only", action="store_true")
    parser.add_argument("--require-strict-highwater", action="store_true")
    parser.add_argument("--batch-baseline", default="", help="本轮开始前已保存名单；仅导出新增，不修改累计名单")
    parser.add_argument("--deliver-current", action="store_true", help="交付已保存名单，不进行采集")
    # ROI 看板参数（可选，用于成本与速率估算）
    parser.add_argument("--budget", type=float, default=0, help="任务预算（用于成本估算）")
    parser.add_argument("--cost-per-reveal", type=float, default=0, help="单次联系方式揭示的估算成本")
    parser.add_argument("--elapsed-minutes", type=float, default=0, help="本次采集耗时（用于速率）")
    args = parser.parse_args()
    if args.deliver_current:
        os.environ["AIPR_JEV_SAVED_ONLY"] = "1"

    source = Path(args.input).resolve()
    payload = json.loads(source.read_text(encoding="utf-8"))
    strategy = json.loads(Path(args.strategy).resolve().read_text(encoding="utf-8"))
    scope_path = Path(args.strategy).resolve().parent / "delivery-exclusion-scope.json"
    if scope_path.exists():
        strategy["deliveryExcludeIdentities"] = json.loads(scope_path.read_text(encoding="utf-8"))["excludeIdentities"]
    if "deliveryExcludeIdentities" in strategy:
        strategy["excludeIdentities"] = list(strategy["deliveryExcludeIdentities"])
    rules = {
        **strategy,
        "threshold": strategy.get("threshold", 78),
        "exclusions": strategy.get("exclusions") or [],
    }
    target_count = max(1, int(strategy.get("targetCount") or 1))
    if args.deliver_current:
        if not args.require_strict_highwater or not payload.get("candidates"):
            raise ValueError("current_delivery_requires_nonempty_strict_highwater")
        target_count = len(payload["candidates"])
    if args.require_strict_highwater:
        audit = payload.get("audit") if isinstance(payload.get("audit"), dict) else {}
        strict_ready = (
            (payload.get("status") == "complete" or (args.deliver_current and payload.get("status") == "incomplete"))
            and payload.get("contact_dedup_mode") == "strict-any-plaintext-value"
            and int(payload.get("strict_selected_count") or 0) >= target_count
            and audit.get("identity_unique") is True
            and audit.get("contact_values_unique") is True
            and int(audit.get("cross_round_contact_overlap") or 0) == 0
        )
        if not strict_ready:
            emit({
                "status": "delivery_incomplete",
                "message": "严格高水位尚未达到交付条件",
                "qualified_count": int(payload.get("strict_selected_count") or 0),
                "target_count": target_count,
                "strict_audit": audit,
            })
            raise SystemExit(7)
    from contact_corrections import read_corrections, apply_corrections
    corrections = read_corrections(os.environ.get("AIPR_CONTACT_CORRECTIONS", ""), args.task_id)
    candidates = apply_corrections(list(payload.get("candidates") or []), corrections)
    if not args.require_strict_highwater:
        saved_path = Path(args.out_dir).resolve() / "aipr_strict_contact_highwater.json"
        if saved_path.exists():
            try:
                saved = json.loads(saved_path.read_text(encoding="utf-8"))
                # Reapply the current delivery rules and contact deduplication to
                # accumulated results as well as this round's contact output.
                candidates = list(saved.get("candidates") or []) + candidates
            except (OSError, ValueError, TypeError):
                pass
    candidates = apply_corrections(candidates, corrections)
    baseline_count = 0
    if args.batch_baseline:
        if not args.deliver_current or not args.require_strict_highwater:
            raise ValueError("batch_scope_requires_saved_strict_delivery")
        baseline = json.loads(Path(args.batch_baseline).read_text(encoding="utf-8"))
        baseline_rows = apply_corrections(baseline.get("candidates") or baseline.get("rows") or [], corrections)
        baseline_count = len(baseline_rows)
        candidates = batch_new_candidates(candidates, baseline_rows)
        if not candidates:
            raise ValueError("batch_has_no_new_qualified_creators")
        target_count = len(candidates)
        rules["excludeContacts"] = list(set(rules.get("excludeContacts") or []) | set().union(*(contact_identity_values(row) for row in baseline_rows)))
    qualified = select_delivery_candidates(candidates, rules, target_count)
    if not args.require_strict_highwater:
        cumulative = {
            "status": "complete" if len(qualified) >= target_count else "incomplete",
            "generated_at": datetime.now().astimezone().isoformat(timespec="seconds"),
            "target_count": target_count, "strict_selected_count": len(qualified),
            "candidate_count": len(qualified), "source_candidate_count": len(candidates),
            "contact_dedup_mode": "strict-any-plaintext-value",
            "audit": _assert_strict_invariants(qualified, rules),
            "strategy": rules, "candidates": qualified,
        }
        atomic_write_json(Path(args.out_dir).resolve() / "aipr_strict_contact_highwater.json", cumulative)
        emit({"status": "strict_contact_audit_ready", "strict_selected_count": len(qualified),
              "target_count": target_count, "message": "已保存累计合格名单，未满目标时继续补采"})
    if len(qualified) < target_count:
        emit({
            "status": "delivery_incomplete",
            "message": f"最终合规且有明文联系方式的达人不足：{len(qualified)}/{target_count}",
            "qualified_count": len(qualified),
            "target_count": target_count,
        })
        raise SystemExit(7)
    flow_path = Path(args.out_dir).resolve() / "aipr_realtime_creator_flow.json"
    if flow_path.exists() and not corrections and not args.batch_baseline:
        from realtime_creator_flow import RealtimeCreatorFlowStore
        RealtimeCreatorFlowStore(flow_path).reconcile_strict_selection({
            "contact_dedup_mode": "strict-any-plaintext-value",
            "target_count": target_count,
            "audit": _assert_strict_invariants(qualified, rules),
            # Selection reruns admission gates; legacy rows may lack this flag.
            "strategy": rules,
            "candidates": [{**row, "precontact_qualified": True} for row in qualified],
        }, finish_contacts=True)
    delivery = build_delivery(qualified, rules, args.task_id)
    from delivery_template import apply_template
    delivery = apply_template(delivery, strategy)
    delivery["contact_corrections_revision"] = len(corrections)
    delivery["source"] = str(source)
    delivery["strategy"] = strategy
    delivery["closed_at_current_count"] = bool(args.deliver_current)
    delivery["planned_target_count"] = int(strategy.get("targetCount") or 1)
    delivery["strict_selected_count"] = len(qualified) if args.batch_baseline else int(payload.get("strict_selected_count") or len(qualified))
    if args.batch_baseline:
        delivery["batch_scope"] = {"baseline_count": baseline_count, "new_count": len(qualified), "cumulative_strict_count": payload.get("strict_selected_count"), "baseline_source": str(Path(args.batch_baseline).resolve())}
    delivery["strict_audit"] = _assert_strict_invariants(qualified, rules)
    delivery["delivery_mode"] = "contacts-only" if args.contacts_only else "robot-handoff"
    if any(row.get("ui_contact_channels_checked_at") for row in qualified):
        baseline_path = Path(args.out_dir).resolve() / "aipr_contact_channels_baseline.json"
        if not baseline_path.is_file():
            baseline_path = source.parent / baseline_path.name
        delivery["contact_channel_completion"] = contact_completion_summary(qualified, target_count, baseline_path)

    # ROI 转化看板：对完整候选集（含未入选）计算漏斗，
    # 这样能看到从候选到交付的真实转化率，而不只是最终名单的计数。
    try:
        from roi_dashboard import build_export_roi
        all_candidates = payload.get("candidates") or []
        elapsed = float(getattr(args, "elapsed_minutes", 0) or 0)
        delivery["roi"] = build_export_roi(all_candidates, qualified, {
            "id": args.task_id,
            "name": args.task_name,
            "budget": float(getattr(args, "budget", 0) or 0),
            "cost_per_reveal": float(getattr(args, "cost_per_reveal", 0) or 0),
            "elapsed_minutes": elapsed,
        })
    except Exception as exc:  # ROI 失败不应阻断交付
        delivery["roi"] = {"error": str(exc)[:200]}

    robot_records = [] if args.contacts_only else list(delivery.get("robot_queue") or [])
    if args.contacts_only:
        delivery["robot_queue"] = []
        delivery["robot_queue_count"] = 0

    output_dir = Path(args.out_dir).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    delivery_path = output_dir / f"aipr_final_delivery_{stamp}.json"
    xlsx_path = output_dir / f"{args.task_name}_最终达人名单_{stamp}.xlsx"
    queue_path = None if args.contacts_only else output_dir / f"{args.task_name}_机器人队列_{stamp}.ndjson"
    batch_path = None if args.contacts_only else output_dir / f"{args.task_name}_机器人批量包_{stamp}.json"
    manifest_path = output_dir / f"{args.task_name}_交接清单_{stamp}.json"
    delivery_path.write_text(json.dumps(delivery, ensure_ascii=False, indent=2), encoding="utf-8")
    write_standard_workbook(delivery["rows"], xlsx_path)
    if queue_path and batch_path:
        queue_path.write_text("".join(json.dumps(row, ensure_ascii=False) + "\n" for row in robot_records), encoding="utf-8")
        batch_path.write_text(json.dumps({
            "schemaVersion": "aipr.robot.batch.v1",
            "contractVersion": delivery["robot_contract_version"],
            "createdAt": delivery["created_at"],
            "task": {"id": args.task_id, "name": args.task_name},
            "dryRunOnly": True,
            "records": robot_records,
        }, ensure_ascii=False, indent=2), encoding="utf-8")
    if args.robot_queue and not args.contacts_only:
        merge_robot_queue(Path(args.robot_queue).resolve(), robot_records)

    original_xlsx = ""
    original = Path(args.original_workbook).resolve() if args.original_workbook else None
    if original and original.exists():
        before = set(output_dir.glob("*.xlsx"))
        exporter = Path(__file__).with_name("export_original_format_with_contacts.py")
        process = subprocess.run([
            sys.executable, str(exporter),
            "--source", str(original),
            "--delivery", str(delivery_path),
            "--out-dir", str(output_dir),
            "--task-name", args.task_name,
        ], text=True, encoding="utf-8", errors="replace", capture_output=True)
        if process.returncode:
            emit({"status": "original_export_error", "message": process.stderr[-800:]})
        else:
            created = sorted(set(output_dir.glob("*.xlsx")) - before, key=lambda path: path.stat().st_mtime, reverse=True)
            original_xlsx = str(created[0]) if created else ""

    artifacts = {
        "finalJson": str(delivery_path),
        "standardXlsx": str(xlsx_path),
        "originalXlsx": original_xlsx,
        "robotNdjson": str(queue_path) if queue_path else "",
        "robotBatchJson": str(batch_path) if batch_path else "",
        "handoffManifest": str(manifest_path),
    }
    delivery["artifacts"] = artifacts
    delivery_path.write_text(json.dumps(delivery, ensure_ascii=False, indent=2), encoding="utf-8")
    manifest = {
        "schemaVersion": "aipr.delivery.handoff.v1",
        "createdAt": delivery["created_at"],
        "task": {"id": args.task_id, "name": args.task_name},
        "outputDirectory": str(output_dir),
        "summary": {
            "candidateCount": delivery["candidate_count"],
            "recommendedCount": delivery["recommended_count"],
            "plainContactCount": delivery["plain_contact_count"],
            "wechatCount": delivery["wechat_contact_count"],
            "phoneCount": delivery["phone_contact_count"],
            "robotRecordCount": len(robot_records),
        },
        # ROI 摘要：交付方最关心的转化率与成本
        "roi": _manifest_roi(delivery.get("roi")),
        "robot": {
            "contractVersion": delivery["robot_contract_version"],
            "modes": [] if args.contacts_only else ["ndjson", "batch-json", "ipc"],
            "dryRunOnly": True,
            "realMessagesSent": 0,
        },
        "artifacts": artifacts,
        "contactChannelCompletion": delivery.get("contact_channel_completion"),
    }
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")

    emit({
        "status": "delivery_ready",
        "output": str(delivery_path),
        "xlsx": str(xlsx_path),
        "original_xlsx": original_xlsx,
        "robot_queue": str(queue_path) if queue_path else "",
        "robot_batch_json": str(batch_path) if batch_path else "",
        "handoff_manifest": str(manifest_path),
        "candidate_count": delivery["candidate_count"],
        "recommended_count": delivery["recommended_count"],
        "plain_contact_count": delivery["plain_contact_count"],
        "wechat_contact_count": delivery["wechat_contact_count"],
        "phone_contact_count": delivery["phone_contact_count"],
        "robot_queue_count": len(robot_records),
    })


if __name__ == "__main__":
    main()
