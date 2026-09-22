from __future__ import annotations

import argparse
import json
from datetime import datetime
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill


def plain(row: dict[str, Any], *keys: str) -> str:
    for key in keys:
        value = str(row.get(key) or "").strip()
        if value:
            return value
    return ""


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--out-dir", required=True)
    args = parser.parse_args()

    source = Path(args.input).resolve()
    payload = json.loads(source.read_text(encoding="utf-8"))
    candidates = payload.get("candidates") or []

    headers = [
        "序号", "达人昵称", "抖音号", "达人等级", "主推类目", "带货形式",
        "近30天销售额下限", "内衣单品", "内衣单品近30天销售额下限",
        "精选联盟主页", "抖音主页", "明文联系方式", "微信", "手机号",
        "内容复核", "联系方式来源", "分跑店铺",
    ]
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "奥斯曼娜达人名单"
    sheet.append(headers)

    for index, row in enumerate(candidates, start=1):
        evaluation = row.get("osmana_evaluation") or {}
        product = evaluation.get("underwear_product_evidence") or {}
        wechat = plain(row, "buyin_contact_wechat", "cart_contact_wechat", "wechat", "微信")
        phone = plain(row, "buyin_contact_phone", "cart_contact_phone", "phone", "手机号")
        contact = plain(row, "plain_contact", "buyin_plain_contact", "cart_plain_contact", "明文联系方式")
        if not contact:
            contact = "；".join(value for value in [wechat, phone] if value)
        sheet.append([
            index,
            plain(row, "nickname", "达人昵称"),
            plain(row, "douyin_id", "unique_id", "抖音号"),
            f"LV{row.get('author_level') or row.get('talent_level') or evaluation.get('level') or ''}",
            plain(row, "category", "主推类目"),
            plain(row, "main_sale_type", "带货形式"),
            row.get("video_sales_low") or row.get("monthly_sales_low") or "",
            evaluation.get("underwear_product_title") or product.get("title") or "",
            evaluation.get("underwear_product_sales_low") or product.get("sales_low") or "",
            plain(row, "buyin_profile_url", "精选联盟主页"),
            plain(row, "douyin_homepage", "抖音主页"),
            contact,
            wechat,
            phone,
            plain(row, "visual_review_note", "content_review_note", "内容复核") or "真人出镜/类目/销量已复核",
            plain(row, "buyin_contact_source", "cart_contact_source", "联系方式来源"),
            plain(row, "shop", "lip_shop"),
        ])

    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="167D68")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    widths = [8, 22, 20, 10, 20, 14, 18, 44, 22, 45, 45, 28, 22, 18, 28, 20, 10]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[openpyxl.utils.get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    out_dir = Path(args.out_dir).resolve()
    out_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output = out_dir / f"奥斯曼娜_当前有效联系方式达人名单_{len(candidates)}人_{stamp}.xlsx"
    workbook.save(output)
    print(json.dumps({"output": str(output), "candidate_count": len(candidates)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
